"""Exportação para Excel: relatório gerencial (mesma organização e nomes da tela) e planilhas-modelo de importação."""
import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from . import rotulos
from .importacao import CABECALHO_MODELO
from .rotulos import COL_VISAO, NOTAS_VISAO

F = 'Segoe UI'
HF = Font(name=F, bold=True, color='FFFFFF', size=10)
HFILL = PatternFill('solid', fgColor='2A5FA8')
NF = Font(name=F, size=10)
BF = Font(name=F, size=10, bold=True)
TOTAL_FILL = PatternFill('solid', fgColor='EAF2FD')
AMARELO = PatternFill('solid', fgColor='FFF4D6')
VERMELHO = PatternFill('solid', fgColor='FBE9E5')
LINHA = Border(bottom=Side(style='thin', color='E1E5EA'))
PCT = '0.0%'
SERIE = ('2A78D6', 'EB6834', '1BAF7A')


def _cabecalho(ws, colunas, linha=1, col0=1):
    for j, (titulo, _, larg, _) in enumerate(colunas, col0):
        c = ws.cell(linha, j, titulo)
        c.font, c.fill = HF, HFILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions[L(j)].width = max(larg + 2, 9)
    ws.row_dimensions[linha].height = 34


def _escala(ws, rng):
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type='num', start_value=0, start_color='F6C7BC',
                                                      mid_type='num', mid_value=0.5, mid_color='FCEBB6',
                                                      end_type='num', end_value=1, end_color='BFE3C8'))


def _tabela(ws, colunas, linhas, linha0=1, col0=1, filtro=True):
    _cabecalho(ws, colunas, linha0, col0)
    for i, r in enumerate(linhas, linha0 + 1):
        for j, (_, k, _, p) in enumerate(colunas, col0):
            c = ws.cell(i, j, r.get(k))
            c.font, c.border = NF, LINHA
            if p == 1:
                c.number_format = PCT
            elif p == 2:
                c.alignment = Alignment(horizontal='center')
    if linhas:
        for j, (_, _, _, p) in enumerate(colunas, col0):
            if p == 1:
                _escala(ws, f'{L(j)}{linha0 + 1}:{L(j)}{linha0 + len(linhas)}')
        if filtro:
            ws.auto_filter.ref = f'{L(col0)}{linha0}:{L(col0 + len(colunas) - 1)}{linha0 + len(linhas)}'
    ws.freeze_panes = ws.cell(linha0 + 1, col0 + 1)
    return linha0 + len(linhas)


def _realcar(ws, linhas, coluna, cores, linha0=1, ncols=None):
    """Pinta a linha inteira conforme o texto da coluna de situação."""
    for i, r in enumerate(linhas, linha0 + 1):
        fill = cores.get(r.get(coluna))
        if fill:
            for j in range(1, (ncols or ws.max_column) + 1):
                ws.cell(i, j).fill = fill


def _grafico(ws, titulo, ref_linhas, colunas_pct, ancora, cab_linha):
    """Colunas agrupadas por unidade usando as colunas % da tabela da visão gerencial."""
    g = BarChart()
    g.type, g.grouping, g.title = 'col', 'clustered', titulo
    g.y_axis.number_format, g.y_axis.scaling.min, g.y_axis.scaling.max = '0%', 0, 1
    g.y_axis.majorGridlines = None
    g.height, g.width, g.gapWidth = 7.5, 16, 80
    a, b = ref_linhas
    for i, col in enumerate(colunas_pct):
        dados = Reference(ws, min_col=col, min_row=cab_linha, max_row=b)
        g.add_data(dados, titles_from_data=True)
        g.series[-1].graphicalProperties.solidFill = SERIE[i]
        g.series[-1].graphicalProperties.line.solidFill = SERIE[i]
    g.set_categories(Reference(ws, min_col=1, min_row=a, max_row=b))
    g.legend.position = 'b'
    ws.add_chart(g, ancora)


def exportar_relatorio(R, destino, data_referencia='', historico=(), criterios=(), incluir_sem_fvi=True):
    rotulos.preparar(R)
    wb = Workbook()

    # ---------------------------------------------------------------- Visão Gerencial
    ws = wb.active
    ws.title = 'Visão Gerencial'
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, 'Visão Gerencial – Avanço do Comissionamento por Unidade – EPC 12').font = Font(name=F, bold=True, size=15)
    ws.cell(2, 1, f'Data de referência: {data_referencia or "–"}   ·   Gerado em {datetime.datetime.now():%d/%m/%Y %H:%M}'
                  f'   ·   TAGs sem FVI gerada entram no cálculo do FVI: {"Sim" if incluir_sem_fvi else "Não"}').font = Font(
        name=F, size=10, color='5F6670')
    # quadro de indicadores gerais
    ws.cell(4, 1, 'Indicadores gerais').font = Font(name=F, bold=True, size=12)
    for j, (k, titulo) in enumerate(rotulos.INDICADORES, 1):
        c = ws.cell(5, j, titulo)
        c.font, c.fill, c.alignment = HF, HFILL, Alignment(horizontal='center', vertical='center', wrap_text=True)
        v = ws.cell(6, j, R.total[k])
        v.number_format, v.font, v.alignment = PCT, Font(name=F, bold=True, size=14), Alignment(horizontal='center')
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 26
    ws.cell(8, 1, 'Avanço por unidade').font = Font(name=F, bold=True, size=12)
    cab = 9
    fim = _tabela(ws, COL_VISAO, R.visao + [R.total], linha0=cab, filtro=False)
    for j in range(1, len(COL_VISAO) + 1):
        ws.cell(fim, j).font, ws.cell(fim, j).fill = BF, TOTAL_FILL
    for i, t in enumerate(NOTAS_VISAO, fim + 2):
        ws.cell(i, 1, t).font = Font(name=F, italic=True, size=9, color='5F6670')
    # gráficos (unidades, sem a linha TOTAL)
    linhas_un = (cab + 1, fim - 1)
    col = {k: j for j, (_, k, _, _) in enumerate(COL_VISAO, 1)}
    g0 = fim + 6
    _grafico(ws, 'FVI por unidade', linhas_un, [col['pct_cii'], col['pct_ciii'], col['pct_cv']], f'A{g0}', cab)
    _grafico(ws, 'FVM por unidade', linhas_un, [col['pct_fvm_ciii'], col['pct_fvm_cv']], f'J{g0}', cab)
    _grafico(ws, 'CCM e TAP por unidade', linhas_un, [col['pct_ccm'], col['pct_tap']], f'A{g0 + 17}', cab)
    _grafico(ws, 'TTAS por unidade', linhas_un, [col['pct_ttas1'], col['pct_ttas2']], f'J{g0 + 17}', cab)
    ws.freeze_panes = None

    # ---------------------------------------------------------------- FVI e FVM (item a item)
    cores = {'Sem FVI gerada': AMARELO, 'Sem SSOP': VERMELHO, 'SSOP fora da base': VERMELHO}
    ws = wb.create_sheet('FVI')
    _tabela(ws, rotulos.COL_TAGS_XLS, R.tags)
    _realcar(ws, R.tags, 'txt_situacao', cores, ncols=len(rotulos.COL_TAGS_XLS))
    ws = wb.create_sheet('FVM')
    _tabela(ws, rotulos.COL_MALHAS_XLS, R.malhas)
    _realcar(ws, R.malhas, 'txt_situacao', cores, ncols=len(rotulos.COL_MALHAS_XLS))

    # ---------------------------------------------------------------- documentos
    for nome, cols, linhas in (('CCM', rotulos.COL_CCM, R.ssop), ('TAP', rotulos.COL_TAP, R.ssop),
                               ('TTAS', rotulos.COL_TTAS, R.sop)):
        _tabela(wb.create_sheet(nome), cols, linhas)

    # ---------------------------------------------------------------- controle
    ws = wb.create_sheet('Verificação')
    ver = [{'g': g, 't': t, 'e': '-' if e is None else e, 'a': a, 'r': {'OK': 'OK', 'VERIFICAR': 'Verificar',
                                                                     'INFO': 'Informação'}[r]}
           for g, t, e, a, r in R.verificacao]
    _tabela(ws, rotulos.COL_VER, ver)
    _realcar(ws, ver, 'r', {'Verificar': VERMELHO}, ncols=len(rotulos.COL_VER))
    _tabela(wb.create_sheet('Pendências'), rotulos.COL_PEND, [dict(zip('dlisp', p)) for p in R.pendencias])
    if historico:
        _tabela(wb.create_sheet('Histórico'), rotulos.COL_HIST, list(historico))
    if criterios:
        ws = wb.create_sheet('Critérios')
        crit = [{'etapa': e, 'codigo': c, 'origem': o} for _, e, c, o in criterios]
        _tabela(ws, rotulos.COL_CRIT, crit)
        for i, r in enumerate(crit, 2):
            if str(r['origem']).startswith('A CONFIRMAR'):
                for j in (1, 2, 3):
                    ws.cell(i, j).fill = AMARELO
    wb.save(destino)


def exportar_modelos(destino):
    """Uma aba por documento com o cabeçalho que o importador reconhece (layout provisório)."""
    wb = Workbook()
    wb.remove(wb.active)
    for tipo, cab in CABECALHO_MODELO.items():
        ws = wb.create_sheet(tipo)
        for j, t in enumerate(cab, 1):
            c = ws.cell(1, j, t)
            c.font, c.fill = HF, HFILL
            ws.column_dimensions[L(j)].width = 26
        ws.freeze_panes = 'A2'
    wb.save(destino)
