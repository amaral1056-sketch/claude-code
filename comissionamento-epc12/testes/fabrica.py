"""Gera planilhas fictícias (mesmo layout das fontes reais) para os testes. Nenhum dado real."""
import openpyxl

from comissionamento.importacao import CADASTRO_ESPERADO, N_COLS_CADASTRO

UN1, UN2 = 'U-9001', 'U-9002'
P1, P2, P3 = 'SOP-U-9001-UT-1001', 'SOP-U-9001-AX-2001', 'SOP-U-9002-EL-7001'
S1, S2, S3, S4 = P1.replace('SOP', 'SSOP') + '-0001', P2.replace('SOP', 'SSOP') + '-0001', \
    P2.replace('SOP', 'SSOP') + '-0002', P3.replace('SOP', 'SSOP') + '-0001'
CT = '0000.0000000.00.0'


def li(caminho, linhas=None):
    linhas = linhas or [(UN1, P1, 'UTILIDADE 1', S1, 'SUB 1'), (UN1, P2, 'AUX 2', S2, 'SUB 2A'),
                        (UN1, P2, 'AUX 2', S3, 'SUB 2B'), (UN2, P3, 'ELÉTRICA 3', S4, 'SUB 3')]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'LISTA'
    ws.append(['LISTA DE INTERFACES – TESTE'])
    ws.append([])
    ws.append(['ITEM', 'UNIDADE', 'SOP', 'DESCRIÇÃO SOP', 'SSOP', 'DESCRIÇÃO SSOP', 'DOCUMENTOS DE REFERÊNCIA'])
    for i, (u, sop, dsop, ssop, dssop) in enumerate(linhas, 1):
        ws.append([i, u, sop, dsop, ssop, dssop, 'DE-0000.00-0000-000-XXX-001   DE-0000.00-0000-000-XXX-002'])
    wb.save(caminho)
    return caminho


def cadastro(caminho, tags, emissao='28/10/2026'):
    """tags: (TAG, SSOP, Nº FVI, STATUS FVI, CII, CIII, CIV, CV[, CLASSE]) – layout do Relatório de Cadastro Geral."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 2, 'Relatório de Cadastro Geral')
    for col in range(1, N_COLS_CADASTRO + 1):
        ws.cell(5, col, CADASTRO_ESPERADO.get(col, f'COLUNA {col}'))
    r = 6
    for t in tags:
        tag, ss, nfvi, aj, ak, al, am, an = t[:8]
        ws.cell(r, 1, tag)
        ws.cell(r, 5, ss)
        ws.cell(r, 6, t[8] if len(t) > 8 else '(032) TRANSFORMADOR')
        ws.cell(r, 29, 'CONTRATO' if tag == 'TAG' else CT)
        ws.cell(r, 31, nfvi)
        for col, v in zip(range(36, 41), (aj, ak, al, am, an)):
            ws.cell(r, col, v)
        r += 1
    ws.cell(r + 1, 1, 'Emissão: ')
    ws.cell(r + 1, 2, 'Classificação da Informação: TESTE')
    if emissao:
        ws.cell(r + 2, 1, f'{emissao}\n\t\t\t17:12:50 PM')
    wb.save(caminho)
    return caminho


def provisorio(caminho, cabecalho, linhas, titulo=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    if titulo:
        ws.append([titulo])
        ws.append([])
    ws.append(cabecalho)
    for r in linhas:
        ws.append(list(r))
    wb.save(caminho)
    return caminho


# casos de borda do teste_sintetico.py original (respostas calculadas à mão)
TAGS_SINTETICAS = [
    ('TAG', 'SSOP', 'Nº FVI', 'STATUS FVI', 'STATUS CAMPO II', 'STATUS CAMPO III', 'STATUS CAMPO IV', 'STATUS CAMPO V'),  # cabeçalho repetido
    ('T1', S1, '1', 'ACAMPO5', 'APROVADO', 'ACAMPO5', 'ACAMPO3E4', 'AAPROV'),       # Campo IV ativo -> teste avisa
    ('T2', S1, '2', 'APROVADO', 'APROVADO', 'APROVADO', 'N/A', 'APROVADO'),
    ('T3', S1, '3', 'ALIBTFUNC', 'N/A', 'ALIBTFUNC', 'N/A', 'N/A'),
    ('T4', S1, '4', 'ACAMPO3E4', 'AAPROV', 'ACAMPO3E4', 'N/A', ''),                 # CV vazio: aplicável, não aprovado
    ('T5', S1, '5', 'CANCELADO', 'CANCELADO', 'CANCELADO', 'N/A', 'CANCELADO'),     # cancelada: fora de tudo
    ('T6', S1, '', '', '', '', '', ''),                                             # sem FVI gerada
    ('T7', ' ' + S2 + ' ', '7', 'ACAMPO3E4', 'APROVADO', ' ALIBTFUNC ', 'N/A', 'XPTO'),  # espaços + código desconhecido
    ('T8', S3, '8', 'APROVADO', 'APROVADO', 'APROVADO', 'N/A', 'APROVADO'),
    ('T9', 'SSOP-NAO-EXISTE', '9', 'APROVADO', 'APROVADO', 'APROVADO', 'N/A', 'APROVADO'),  # fora da base
    ('T10', '', '10', 'ARECEB', 'ARECEB', 'ARECEB', 'N/A', 'ARECEB'),                       # sem SSOP (não alocada)
]
FVM_SINTETICO = [('M1', P1, S1, 'ACAMPO5'), ('M2', P1, S1, 'APROVADO'), ('M3', P1, S1, 'AAPRV3E4FIS'),
                 ('M4', P1, S1, 'CANCEL'), ('M5', P1, S1, 'LIBTFUNC')]
CCM_SINTETICO = [('C1', S1, 'CANCELADO'), ('C1R1', S1, 'CCMCONCLUIDO'), ('C2', S2, 'AGAPROVA')]
# TAP por SOP, como a lista da FIC: Id | Número do TAP ("TAP-" + SOP) | Descrição | Status
TAP_SINTETICO = [('1', 'TAP-' + P1, 'UTILIDADE 1', 'TAP CONCLUIDO'), ('2', 'TAP-' + P2, 'AUX 2', 'TAP1APROVCQ'),
                 ('3', 'TAP-' + P2, 'AUX 2', 'XYZ')]
CABECALHO_TAP = ['Id FIC', 'Número do TAP', 'Descrição', 'Status']
TTAS_SINTETICO = [('TT1', P1, 'TTASCONCLUIDO'), ('TT2', P2, 'TTAS1APROVADO')]


def cadastros_basicos(caminho, linhas):
    """Lista de Cadastros Básicos da FIC (base SOP / SSOP): (Código, Descrição, Tipo)."""
    return provisorio(caminho, ['Código', 'Descrição', 'Tipo', 'Status', 'Disciplina', 'Empreendimento'],
                      [(c, d, t, 'OPERAÇÃO', '', 'EMPREENDIMENTO TESTE') for c, d, t in linhas])


BASICOS = [(P1, 'UTILIDADE 1', 'SOP'), (P2, 'AUX 2', 'SOP'), (P3, 'ELÉTRICA 3', 'SOP'),
           (S1, 'SUB 1', 'SSOP'), (S2, 'SUB 2A', 'SSOP'), (S3, 'SUB 2B', 'SSOP'), (S4, 'SUB 3', 'SSOP')]


def hierarquia_malhas(caminho, malhas, emissao='29/10/2026'):
    """Relatório de Hierarquia de Malhas: malhas = [(malha, status, [(TIPO, tag, SSOP, SOP), ...])]."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 6, 'Relatório de Hierarquia de Malhas')
    ws.append([])
    ws.append([])
    ws.append(['Empreendimento:', None, None, None, None, 'EMPREENDIMENTO TESTE'])
    ws.append([])
    ws.append(['TIPO', 'MALHA', 'TAG/LINHAS', 'STATUS', 'DESCRIÇÃO', 'SSOP', 'SOP', 'UNIDADE', 'DISCIPLINA', 'CLASSE'])
    for malha, status, filhos in malhas:
        ws.append(['MALHA', malha, '-', status, 'MALHA DE TESTE', '-', '-', '-', 'INS', '-'])
        for tipo, tag, ssop, sop in filhos:
            ws.append([tipo, malha, tag, 'OPERAÇÃO', 'ITEM', ssop, sop, 'UNIDADE TESTE', 'INS', '(000) TESTE'])
    ws.append(['Emissão: ', 'USUÁRIO TESTE', 'Classificação da Informação: TESTE', 'Folha: 1 de 1'])
    ws.append([f'{emissao}  09:00:19'])
    wb.save(caminho)
    return caminho


def status_malhas(caminho, linhas, filtro='Status FVM: Aprovado, Cancelado, ', emissao='29/10/2026'):
    """Relatório de Status de Malhas: linhas = [(malha, Nº FVM, status)] – título, filtros, cabeçalho, dados, rodapé."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 5, 'Relatório de Status de Malhas')
    ws.cell(4, 1, 'Listagem de filtros selecionados:')
    ws.cell(5, 1, 'Disciplina: INS')
    ws.cell(8, 1, filtro)
    ws.cell(10, 6, 'Status de Atividades - FVM')
    for j, t in enumerate(['Disciplina', 'Malha', 'Descrição', 'FVM', 'Contratada', 'Status', 'No Prazo?'], 1):
        ws.cell(11, j, t)
    r = 13
    for malha, n_fvm, st in linhas:
        for j, v in enumerate(['INS', malha, 'MALHA DE TESTE', n_fvm or None, 'CONTRATADA TESTE', st or None,
                               'Data Prevista não Disponível'], 1):
            ws.cell(r, j, v)
        r += 1
    for j, v in enumerate(['Emissão: ', 'USUÁRIO TESTE', None, 'Classificação da Informação: TESTE', None, 'Folha: 1 de 1'], 1):
        ws.cell(r, j, v)
    ws.cell(r + 1, 1, f'{emissao}\n\t\t\t09:42:54 AM')
    wb.save(caminho)
    return caminho


def lista_fvm(caminho, linhas):
    """Lista de FVM da FIC: [(Código FVM, malha, status – código)]."""
    return provisorio(caminho, ['Código FVM', 'Descrição', 'Malha', 'Disciplina', 'Código do Template', 'Status',
                                'Empreendimento', 'Rotear', 'Grupo de Proprietários'],
                      [(c, f'FVM - {m}-MALHA DE TESTE', m, 'INS', 'FVM-INS001', st, 'EMPREENDIMENTO TESTE', 'N', '')
                       for c, m, st in linhas])


def fvm(pasta, linhas):
    """FVM BASE (hierarquia: uma TAG por malha e SSOP) e a Lista de FVM (um status por malha) a partir de
    [(malha, SOP, SSOP, status)]. Devolve (base, lista)."""
    malhas = {}
    for m, sop, ssop, st in linhas:
        malhas.setdefault(m, [st, []])[1].append(('TAG', f'T-{m}-{len(malhas[m][1]) + 1}', ssop, sop))
    base = hierarquia_malhas(pasta / 'fvm_base.xlsx', [(m, 'OPERAÇÃO', filhos) for m, (_, filhos) in malhas.items()])
    lista = lista_fvm(pasta / 'fvm_lista.xlsx', [(str(100 + i), m, st) for i, (m, (st, _)) in enumerate(malhas.items())])
    return base, lista


def templates_fvm_pdf(caminho, templates):
    """PDF no formato do "Template de FVM" da FIC: templates = [(código, revisão, disciplina, [(item, texto)])].
    Reproduz o texto colado do original ("TESTES FUNCIONAISV.7 ...")."""
    from fpdf import FPDF
    pdf = FPDF(orientation='L')
    pdf.set_font('Helvetica', size=9)
    for codigo, rev, disc, perguntas in templates:
        pdf.add_page()
        for linha in ('Template de FVM - Folha de Verificação de Malhas EMPREENDIMENTO TESTE',
                      f'Template: {codigo}Revisão: {rev}Status: Ativo',
                      f'CÓDIGO DA FVM:MALHA:DISCIPLINA: {disc}DOCS REF .:',
                      'CAMPO III / IV III.1III.1III.2 CAMPO V V.7V.8',
                      ' III / IV) INSPEÇÃO MECÂNICA ' + ' '.join(f'{i} {t}' for i, t in perguntas if not i.startswith('V.')),
                      'ACEITAÇÃO OBSERVAÇÕESATIVIDADE EXECUÇÃO',
                      'V) TESTES FUNCIONAIS' + ' '.join(f'{i} {t}' for i, t in perguntas if i.startswith('V.')),
                      'ACEITAÇÃO OBSERVAÇÕESATIVIDADE EXECUÇÃO',
                      'Emissão: USUÁRIO TESTE Classificação da Informação: TESTE Folha: 1 de 1'):
            pdf.multi_cell(0, 5, linha, new_x='LMARGIN', new_y='NEXT')
    pdf.output(str(caminho))
    return caminho


def templates_fvi_pdf(caminho, templates):
    """PDF no formato do "Template de FVI" da FIC: templates = [(código, classificação, {campo: [perguntas]})].
    Cada campo (II, III, V) começa numa página nova, com o cabeçalho e o rodapé repetidos, como no original;
    perguntas longas quebram em várias linhas."""
    from fpdf import FPDF
    titulos = {'II': 'II) INSPEÇÃO TÉCNICA', 'III': 'III) INSPEÇÃO MECÂNICA', 'V': 'V) TESTES FUNCIONAIS'}
    pdf = FPDF()
    pdf.set_font('Helvetica', size=9)

    def linha(t):
        pdf.multi_cell(0, 5, t, new_x='LMARGIN', new_y='NEXT')
    for codigo, classe, campos in templates:
        for k, (campo, perguntas) in enumerate(campos.items()):
            pdf.add_page()
            for t in ('Template de FVI - Folha de Verificação de Item', 'EMPREENDIMENTO TESTE / 0000.0000000.00.0',
                      f'Template: {codigo}', 'Revisão: 1.0', 'Status: Ativo'):
                linha(t)
            if k == 0:
                for t in ('TAG: SKID:', f'CLASSIFICAÇÃO: {classe} MÓDULO:', 'DISCIPLINA: TESTE RM:',
                          'I) CARACTERÍSTICAS DO ITEM', 'II) DADOS DO RECEBIMENTO', 'DATA RESPONSÁVEL OBSERVAÇÃO'):
                    linha(t)
            linha(titulos[campo])
            for n, p in enumerate(perguntas, 1):
                linha(f'{n}) {p}')
            for t in ('ACEITAÇÃO', 'ATIVIDADE EXECUÇÃO PETROBRAS', 'Emissão: USUÁRIO TESTE Classificação da Informação: TESTE',
                      '28/10/2026  13:39:59 EMPREENDIMENTO TESTE'):
                linha(t)
    pdf.output(str(caminho))
    return caminho
