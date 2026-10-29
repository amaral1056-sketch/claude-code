"""Importadores: layouts aceitos, recusas (nada é gravado) e exportações."""
import datetime

import openpyxl
import pytest

from comissionamento import exportacao, importacao, servico, tema
from comissionamento.banco import Banco
from comissionamento.importacao import ErroImportacao
from testes import fabrica as f


@pytest.fixture
def banco(tmp_path):
    b = Banco(tmp_path / 'teste.db')
    servico.importar_li(b, f.li(tmp_path / 'li.xlsx'))
    yield b
    b.fechar()


def test_li_base_e_documentos(banco):
    un, sops, ssops = banco.base()
    assert un == [f.UN1, f.UN2]
    assert [s['sop'] for s in sops] == [f.P1, f.P2, f.P3]
    assert ssops[0]['grupo'] == 'UT'
    assert ssops[0]['documentos'] == 'DE-0000.00-0000-000-XXX-001 | DE-0000.00-0000-000-XXX-002'


@pytest.mark.parametrize('linhas, msg', [
    ([(f.UN1, f.P1, 'A', f.S1, 'X'), (f.UN1, f.P1, 'A', f.S1, 'Y')], 'SSOP duplicado'),
    ([(f.UN1, f.P1, 'A', f.S2, 'X')], 'não começa com'),
    ([(f.UN2, f.P1, 'A', f.S1, 'X')], 'não confere'),
])
def test_li_com_problema_nao_grava(banco, tmp_path, linhas, msg):
    with pytest.raises(ErroImportacao, match=msg):
        servico.importar_li(banco, f.li(tmp_path / 'ruim.xlsx', linhas))
    assert len(banco.base()[2]) == 4   # base anterior preservada


def test_li_sem_aba_lista(tmp_path):
    p = tmp_path / 'x.xlsx'
    openpyxl.Workbook().save(p)
    with pytest.raises(ErroImportacao, match='LISTA'):
        importacao.ler_li(p)


def test_cadastro_layout_mudou_para(banco, tmp_path):
    p = f.cadastro(tmp_path / 'c.xlsx', f.TAGS_SINTETICAS)
    wb = openpyxl.load_workbook(p)
    wb.active.cell(5, 37, 'CAMPO II')          # coluna AK renomeada
    wb.save(p)
    with pytest.raises(ErroImportacao, match='AK'):
        servico.importar_fic(banco, 'FVI', p)
    assert banco.dados('FVI') == []


def test_cadastro_data_e_regra_de_linha(tmp_path):
    regs, data, info = importacao.ler_cadastro(f.cadastro(tmp_path / 'c.xlsx', f.TAGS_SINTETICAS))
    assert data == datetime.date(2026, 10, 28)
    assert [r['tag'] for r in regs] == [f'T{i}' for i in range(1, 11)]   # sem cabeçalho repetido e sem rodapé
    assert regs[0]['dados']['STATUS CAMPO IV'] == 'ACAMPO3E4'


def test_cadastro_sem_rodape_usa_data_informada(banco, tmp_path):
    p = f.cadastro(tmp_path / 'c.xlsx', f.TAGS_SINTETICAS, emissao=None)
    _, data, _ = servico.importar_fic(banco, 'FVI', p, datetime.date(2026, 11, 1))
    assert data == '2026-11-01'


def test_status_de_malhas_sem_rodape_e_com_filtro(banco, tmp_path):
    p = f.status_malhas(tmp_path / 'fvm.xlsx', [('M1', '10', 'Aprovado'), ('M2', '', '')], emissao='30/10/2026')
    regs, data, info = servico.importar_fic(banco, 'FVM', p)
    assert [(r['malha'], r['n_fvm'], r['status']) for r in regs] == [('M1', '10', 'Aprovado'), ('M2', '', '')]  # sem rodapé
    assert data == '2026-10-30' and info['tipo'] == 'FVM_STATUS'
    assert info['filtro_status'] == 'Status FVM: Aprovado, Cancelado'


def test_provisorio_sem_cabecalho_reconhecido(banco, tmp_path):
    p = f.provisorio(tmp_path / 'x.xlsx', ['A', 'B', 'C'], [('1', '2', '3')])
    with pytest.raises(ErroImportacao, match='Cabeçalho de CCM'):
        servico.importar_fic(banco, 'CCM', p)


def test_status_oficial_pelo_nome_do_arquivo(banco, tmp_path):
    p = f.provisorio(tmp_path / 'STATUS_TAP.xlsx', ['Valor', 'Descrição'], [('NOVO', 'Novo código'), ('TAP CONCLUIDO', 'x')])
    assert servico.importar_status(banco, p) == ('TAP', 2)
    assert [c for d, c, _ in banco.status_fic() if d == 'TAP'] == ['NOVO', 'TAP CONCLUIDO']


def test_relatorio_excel_e_backup(banco, tmp_path):
    servico.importar_fic(banco, 'FVI', f.cadastro(tmp_path / 'c.xlsx', f.TAGS_SINTETICAS))
    R = servico.calcular(banco)
    destino = tmp_path / 'rel.xlsx'
    xlsx, pdf = exportacao.exportar_relatorio(R, destino, '28/10/2026', True, banco.status_fic())
    assert (xlsx, pdf) == (destino, tmp_path / 'rel.pdf')
    wb = openpyxl.load_workbook(destino)
    assert wb.sheetnames == ['Visão Gerencial', 'FVI', 'FVM', 'CCM', 'TAP', 'TTAS']      # só as páginas pedidas
    assert all(len(ws._images) == 1 for ws in wb)                                       # logo em cada aba
    # PDF: Visão Gerencial + uma página por documento, fontes do programa embutidas
    conteudo = pdf.read_bytes()
    assert conteudo.startswith(b'%PDF') and conteudo.count(b'/Type /Page\n') >= 6
    assert b'EPC12Display' in conteudo and b'IBMPlexMono' in conteudo
    vg = wb['Visão Gerencial']
    assert len(vg._charts) == 5                                  # pizzas FVI, FVM, CCM, TAP e TTAS
    assert all(type(g).__name__ == 'PieChart' for g in vg._charts)
    # mesma cor da tela para o mesmo status (APROVADO é a fatia maior do FVI)
    ordem_fvi = [c for d, c, _ in banco.status_fic() if d == 'FVI']
    tipo, cor = tema.cores_pizza(['APROVADO'], ordem_fvi)[0]
    assert tipo == 'solido' and cor in tema.CATEGORICAS                    # paleta categórica colorida
    assert vg._charts[0].series[0].dPt[0].graphicalProperties.solidFill.srgbClr == cor.lstrip('#').upper()
    # cada código da lista oficial com uma cor diferente, na ordem fixa da paleta; N/A e sem status em cinza
    assert tema.cores_pizza(['A', 'B', 'C', 'N/A', 'Sem status'], ['A', 'B', 'C']) == [
        ('solido', tema.CATEGORICAS[0]), ('solido', tema.CATEGORICAS[1]), ('solido', tema.CATEGORICAS[2]),
        tema.PIZZA_NA, tema.PIZZA_SEM]
    # números das pizzas na própria aba, abaixo dos gráficos
    assert [vg['A60'].value, vg['A61'].value, vg['A62'].value, vg['B62'].value] == [
        'FVI – Status FVI por TAG', 'Status (FIC)', 'APROVADO', 3]
    blocos = [c.value for c in vg['A'] if isinstance(c.value, str) and c.value.split(' – ')[0] in
              ('FVI', 'FVM', 'CCM', 'TAP', 'TTAS')]
    assert blocos == ['FVI – Status FVI por TAG', 'FVM – Status da FVM por malha', 'CCM – Status do CCM por SSOP',
                      'TAP – Status do TAP por SOP', 'TTAS – Status do TTAS por SOP']
    fvi = list(wb['FVI'].iter_rows(min_row=4, values_only=True))      # linhas 1-2: logo e título
    assert fvi[0] == ('Unidade', 'SOP', 'SSOP', 'TAG', 'Descrição da TAG', 'Nº da FVI', 'Status FVI', 'FVI Campo II',
                      'FVI Campo III', 'FVI Campo V')
    assert len(fvi) == 11
    t1 = next(r for r in fvi if r[3] == 'T1')
    assert t1[6:] == ('ACAMPO5', 'APROVADO', 'ACAMPO5', 'AAPROV')
    assert [c.value for c in wb['CCM'][4]] == ['Unidade', 'SOP', 'SSOP', 'Descrição da SSOP', 'Status do CCM']
    assert [c.value for c in wb['TTAS'][4]] == ['Unidade', 'SOP', 'Descrição do SOP', 'Situação TTAS 1', 'Situação TTAS 2']
    assert [c.value for c in wb['TAP'][4]] == ['Unidade', 'SOP', 'Descrição do SOP', 'Status do TAP']
    copia = tmp_path / 'copia.db'
    banco.copiar_para(copia)
    b2 = Banco(copia)
    assert len(b2.dados('FVI')) == 10
    b2.fechar()
